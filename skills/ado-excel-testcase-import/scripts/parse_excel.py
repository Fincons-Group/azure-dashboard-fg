"""Parse an ALM-style test-case Excel into cases.json.

Layout expected per sheet (the "Pacchetto END-TO-END" export): a row with col A == "Test Case"
starts a case (B title, I description); following rows with a value in C are steps
(D action, E expected result). Sheets with no "Test Case" rows are skipped.

usage: parse_excel.py <xlsx> [--sheets A,B] [--fill FFFFFF00,...] [--out cases.json] [--colors]
  --colors   only report which fill colours exist and how many cases each touches (use before asking
             the user about a colour filter)
  --fill     keep only cases with at least one cell in that fill colour (ARGB hex)
"""
import argparse, collections, json, sys
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
ap = argparse.ArgumentParser()
ap.add_argument("xlsx")
ap.add_argument("--sheets")
ap.add_argument("--fill")
ap.add_argument("--out", default="cases.json")
ap.add_argument("--colors", action="store_true")
ap.add_argument("--title-fill", action="store_true", help="with --fill: match only the title cell's colour")
a = ap.parse_args()

NEUTRAL = {"FFFFFFFF", "00000000"}


def fill_of(cell):
    f = cell.fill
    if not f or not f.fill_type:
        return None
    c = f.fgColor
    if c.type == "rgb" and c.rgb not in NEUTRAL:
        return c.rgb
    if c.type == "theme" and c.theme not in (0,):
        return f"theme{c.theme}{'+%.2f' % c.tint if c.tint else ''}"
    return None


def priority(n):
    return 1 if n >= 5 else 2 if n >= 3 else 3 if n == 2 else 4


wb = openpyxl.load_workbook(a.xlsx)
wanted = set(a.sheets.split(",")) if a.sheets else None
keep = set(x.strip().upper() for x in a.fill.split(",")) if a.fill else None
cases, colors = [], collections.defaultdict(set)
for ws in wb.worksheets:
    if wanted and ws.title not in wanted:
        continue
    sheet_cases = []
    header = {str(ws.cell(1, c).value or "").strip(): c for c in range(1, ws.max_column + 1)}
    if "Work Item Type" in header and "Title" in header:
        # ADO query export: ID | Work Item Type | Title | Test Step | Step Action | Step Expected | ...
        L = {"type": header["Work Item Type"], "title": header["Title"], "step": header.get("Test Step"),
             "action": header.get("Step Action"), "expected": header.get("Step Expected"),
             "id": header.get("ID"), "ext": header.get("External ID"), "desc": header.get("Description"),
             "first": 2, "cols": range(1, ws.max_column + 1), "area_is_sheet": True}
    else:
        # ALM "Pacchetto" export: A "Test Case" | B title | C step no | D action | E expected | I description
        L = {"type": 1, "title": 2, "step": 3, "action": 4, "expected": 5, "id": None, "ext": None, "desc": 9,
             "first": 1, "cols": range(1, 10), "area_is_sheet": False}
    val = lambda r, c: ws.cell(r, c).value if c else None
    for r in range(L["first"], ws.max_row + 1):
        if val(r, L["type"]) == "Test Case":
            sheet_cases.append({"key": f"{ws.title}!{r}", "sheet": ws.title, "row": r,
                                "title": str(val(r, L["title"]) or "").strip(),
                                "desc": str(val(r, L["desc"]) or "").strip(),
                                "source_id": val(r, L["id"]), "external_id": val(r, L["ext"]),
                                "area_is_sheet": L["area_is_sheet"],
                                "steps": [], "fills": set(), "title_fill": fill_of(ws.cell(r, L["title"]))})
        elif sheet_cases and val(r, L["step"]) not in (None, ""):
            sheet_cases[-1]["steps"].append([val(r, L["action"]), val(r, L["expected"])])
        if sheet_cases:
            for col in L["cols"]:
                f = fill_of(ws.cell(r, col))
                if f:
                    sheet_cases[-1]["fills"].add(f)
                    colors[f].add(sheet_cases[-1]["key"])
    cases += sheet_cases

if a.colors:
    print("fill colour | cases touched")
    for f, ks in sorted(colors.items(), key=lambda kv: -len(kv[1])):
        print(f, "|", len(ks), "|", ", ".join(sorted(ks)[:8]))
    print("total cases:", len(cases), "sheets:", sorted({c["sheet"] for c in cases}))
    sys.exit()

if keep:
    if a.title_fill:
        cases = [c for c in cases if (c["title_fill"] or "").upper() in keep]
    else:
        cases = [c for c in cases if c["fills"] & keep]

# Area = title prefix before " - " (or ":"); unify spelling variants by first word.
# ADO-export titles start with an id code (NF-CENS-004 - ...), so there the tab name is the area.
raw = [c["sheet"] if c["area_is_sheet"] else c["title"].split(" - ")[0].split(":")[0].strip() for c in cases]
by_first = collections.defaultdict(collections.Counter)
for r in raw:
    by_first[r.split()[0].lower() if r else ""][r] += 1
for c, r in zip(cases, raw):
    c["area"] = by_first[r.split()[0].lower() if r else ""].most_common(1)[0][0]
    c["priority"] = priority(len(c["steps"]))
    c["fills"] = sorted(c["fills"])

json.dump(cases, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(f"{len(cases)} cases -> {a.out}")
for s in sorted({c["sheet"] for c in cases}):
    sc = [c for c in cases if c["sheet"] == s]
    print(f"\n[{s}] {len(sc)} cases, {sum(len(c['steps']) for c in sc)} steps")
    for c in sc:
        print(f"  {c['key']} | P{c['priority']} | {len(c['steps'])} | {c['area']} | {c['title']}")
